import json
import datetime
import hashlib
import uuid
import re
import base64
import io
from typing import Optional
from ..utils.gemini_client import GeminiClient
from ..config import *
from ..core.bicameral_engine import BicameralEngine, ReasoningMode, ReasoningRequest

def _create_thought_signature(user_id, agent, prompt, response, db_helper, context=None):
    """Create and store thought signature - shared helper."""
    try:
        thought_id = f"thought_{uuid.uuid4().hex[:12]}"
        timestamp = datetime.datetime.now().isoformat()
        context_hash = hashlib.sha256(f"{user_id}:{prompt[:200]}:{timestamp}".encode()).hexdigest()[:16]
        
        thought_data = {
            "thought_id": thought_id, "timestamp": timestamp, "agent": agent,
            "context_hash": context_hash,
            "reasoning_trace": [f"Prompt: {prompt[:100]}...", f"Response: {response[:200]}..."],
            "confidence": 0.85, "tool_calls": [], "action_output": response[:500], "parent_signature": ""
        }
        
        db_helper.create_thought_signature(user_id, thought_data)
        db_helper.update_agent_memory_full(user_id, thought_sig_dict=thought_data, active_agents=agent, reasoning_mode='DEEP')
        
        if context: context.log(f"🧠 Thought signature stored: {thought_id}")
        return thought_id
    except Exception as e:
        print(f"❌ Thought signature error: {e}")
        return None


# =============================================================================
# PDF TEXT EXTRACTION
# =============================================================================
def extract_text_from_pdf(pdf_base64: str) -> str:
    """
    Extract text from a base64-encoded PDF file.
    Uses pypdf library for extraction.
    """
    try:
        from pypdf import PdfReader
        
        # Decode base64 to bytes
        pdf_bytes = base64.b64decode(pdf_base64)
        pdf_stream = io.BytesIO(pdf_bytes)
        
        # Read PDF and extract text
        reader = PdfReader(pdf_stream)
        text_parts = []
        
        for page in reader.pages:
            page_text = page.extract_text()
            if page_text:
                text_parts.append(page_text)
        
        return "\n\n".join(text_parts)
    except Exception as e:
        print(f"❌ PDF extraction error: {e}")
        return ""


# =============================================================================
# DOCX RESUME GENERATOR
# =============================================================================
def generate_resume_docx(resume_data: dict) -> str:
    """
    Generate a DOCX file from structured resume data.
    Returns base64-encoded DOCX bytes.
    """
    try:
        from docx import Document
        from docx.shared import Inches, Pt
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        
        doc = Document()
        
        # Set narrow margins
        for section in doc.sections:
            section.top_margin = Inches(0.5)
            section.bottom_margin = Inches(0.5)
            section.left_margin = Inches(0.75)
            section.right_margin = Inches(0.75)
        
        resume = resume_data.get("resume", resume_data)
        
        # Title / Name placeholder
        title = doc.add_heading("[YOUR NAME]", 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        # Contact info placeholder
        contact = doc.add_paragraph("[Email] | [Phone] | [LinkedIn] | [GitHub]")
        contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        # Professional Summary
        summary = resume.get("professional_summary", "")
        if summary:
            doc.add_heading("Professional Summary", level=1)
            doc.add_paragraph(summary)
        
        # Skills Section
        skills = resume.get("skills_section", [])
        if skills:
            doc.add_heading("Technical Skills", level=1)
            skills_text = " • ".join(skills)
            doc.add_paragraph(skills_text)
        
        # Experience Section
        experiences = resume.get("experience_section", [])
        if experiences:
            doc.add_heading("Experience", level=1)
            for exp in experiences:
                # Title and Duration
                p = doc.add_paragraph()
                title_run = p.add_run(f"{exp.get('title', 'Role')} | {exp.get('organization', 'Organization')}")
                title_run.bold = True
                p.add_run(f"  |  {exp.get('duration', '')}")
                
                # Bullets
                for bullet in exp.get("bullets", []):
                    doc.add_paragraph(bullet, style='List Bullet')
        
        # Projects Section
        projects = resume.get("projects_section", [])
        if projects:
            doc.add_heading("Projects", level=1)
            for proj in projects:
                p = doc.add_paragraph()
                name_run = p.add_run(proj.get("name", "Project"))
                name_run.bold = True
                
                tech = proj.get("technologies", [])
                if tech:
                    p.add_run(f"  |  {', '.join(tech) if isinstance(tech, list) else tech}")
                
                desc = proj.get("description", "")
                if desc:
                    doc.add_paragraph(desc)
                
                for achievement in proj.get("achievements", []):
                    doc.add_paragraph(achievement, style='List Bullet')
        
        # Education Section
        education = resume.get("education_section", {})
        if education:
            doc.add_heading("Education", level=1)
            p = doc.add_paragraph()
            degree_run = p.add_run(f"{education.get('degree', 'Degree')}")
            degree_run.bold = True
            p.add_run(f"  |  {education.get('institution', 'University')}  |  {education.get('year', '')}")
            
            coursework = education.get("relevant_coursework", [])
            if coursework:
                doc.add_paragraph(f"Relevant Coursework: {', '.join(coursework)}")
        
        # Save to bytes
        docx_stream = io.BytesIO()
        doc.save(docx_stream)
        docx_stream.seek(0)
        
        # Encode as base64
        return base64.b64encode(docx_stream.read()).decode('utf-8')
        
    except Exception as e:
        print(f"❌ DOCX generation error: {e}")
        return ""


# =============================================================================
# ARMORY: ATS RESUME ANALYZER (Enhanced)
# =============================================================================
async def analyze_resume_ats(db_helper, user_id, resume_text: str, job_description: str, context=None) -> dict:
    """
    Advanced ATS Resume Analyzer using Gemini 3 Deep Reasoning.
    Returns comprehensive analysis with score breakdown, improvements, and checklist.
    """
    engine = BicameralEngine()
    
    prompt = f"""You are an elite ATS (Applicant Tracking System) Analyzer and Career Coach.

RESUME TEXT:
{resume_text[:8000]}

JOB DESCRIPTION:
{job_description[:5000]}

═══ YOUR MISSION ═══
Perform a COMPREHENSIVE analysis of this resume against the job description.

═══ ANALYSIS FRAMEWORK ═══

1. **KEYWORD MATCH ANALYSIS**
   - Extract 15-20 key skills/requirements from JD
   - Check which are present/missing in resume
   - Calculate keyword match percentage

2. **EXPERIENCE FIT**
   - Years of experience alignment
   - Role relevance (0-100)
   - Industry context match

3. **FORMAT & STRUCTURE**
   - ATS-friendly formatting check
   - Section completeness (Contact, Summary, Experience, Education, Skills)
   - Bullet point effectiveness

4. **QUANTIFICATION CHECK**
   - Count metrics/numbers in achievements
   - Identify weak bullets that need quantification

5. **ACTION VERBS ANALYSIS**
   - Check for strong action verbs
   - Identify passive language

═══ OUTPUT JSON ═══
Return ONLY valid JSON:
{{
    "ats_score": 78,
    "score_breakdown": {{
        "keyword_match": 72,
        "experience_fit": 85,
        "format_structure": 80,
        "quantification": 65,
        "action_verbs": 90
    }},
    "matched_keywords": ["Python", "Machine Learning", "API Development"],
    "missing_keywords": ["Kubernetes", "AWS", "Agile"],
    "critical_improvements": [
        {{
            "issue": "Missing cloud experience",
            "current": "No cloud platforms mentioned",
            "suggested": "Add AWS/GCP projects or certifications",
            "priority": "HIGH"
        }}
    ],
    "bullet_improvements": [
        {{
            "original": "Worked on backend systems",
            "improved": "Engineered RESTful APIs serving 10K+ daily requests with 99.9% uptime",
            "reason": "Added quantification and specific technology"
        }}
    ],
    "checklist": [
        {{"item": "Add AWS certification", "status": "pending", "impact": "HIGH"}},
        {{"item": "Quantify project impacts", "status": "pending", "impact": "HIGH"}},
        {{"item": "Add skills section with JD keywords", "status": "pending", "impact": "MEDIUM"}}
    ],
    "overall_assessment": "Your resume shows strong technical foundation but lacks cloud experience critical for this role. Focus on AWS/Kubernetes projects.",
    "interview_ready": false,
    "estimated_pass_rate": "60%"
}}
"""

    try:
        response = await engine.reason(ReasoningRequest(
            prompt=prompt,
            user_id=user_id,
            agent="campaign",
            mode=ReasoningMode.DEEP,
            max_thinking_tokens=16384,
            system_instruction="You are an expert ATS analyzer. Output valid JSON only. Be thorough and actionable."
        ))
        
        # Parse JSON from response
        content = response.content
        if "```" in content:
            content = re.sub(r'```(?:json)?\s*', '', content).replace('```', '')
        
        match = re.search(r'\{[\s\S]*\}', content)
        if match:
            result = json.loads(match.group())
        else:
            result = json.loads(content)
        
        # Save to resume history
        db_helper.update_campaign_state(user_id, {
            "last_ats_analysis": {
                "score": result.get('ats_score', 0),
                "job_description_preview": job_description[:200],
                "interview_ready": result.get('interview_ready', False)
            }
        })
        
        # Create thought signature
        _create_thought_signature(
            user_id, "campaign", 
            f"ATS Analysis for JD: {job_description[:100]}", 
            f"Score: {result.get('ats_score', 'N/A')} - {result.get('overall_assessment', '')[:200]}", 
            db_helper, context
        )
        
        return {"success": True, "data": result}
        
    except Exception as e:
        print(f"❌ ATS Analysis Error: {e}")
        return {"success": False, "error": str(e)}


# =============================================================================
# ARMORY: RESUME GENERATOR
# =============================================================================
async def generate_tailored_resume(db_helper, user_id, job_description: str, projects: list, 
                                    skills: Optional[list] = None, education: Optional[str] = None, 
                                    experience: Optional[str] = None, context=None) -> dict:
    """
    Generate a tailored resume based on job description and user's projects.
    Returns structured resume content with suggestions.
    """
    engine = BicameralEngine()
    
    # Build projects context
    projects_text = ""
    for i, proj in enumerate(projects, 1):
        if isinstance(proj, dict):
            projects_text += f"""
PROJECT {i}: {proj.get('name', 'Untitled')}
- Description: {proj.get('description', 'N/A')}
- Technologies: {proj.get('technologies', 'N/A')}
- Impact: {proj.get('impact', 'N/A')}
"""
        else:
            projects_text += f"PROJECT {i}: {proj}\n"
    
    prompt = f"""You are an expert Resume Writer and Career Coach.

═══ JOB DESCRIPTION ═══
{job_description[:5000]}

═══ USER'S PROJECTS ═══
{projects_text[:6000]}

═══ ADDITIONAL CONTEXT ═══
Skills: {skills or 'Not provided - infer from projects'}
Education: {education or 'Not provided'}
Past Experience: {experience or 'Student/Fresh Graduate'}

═══ YOUR MISSION ═══
Create a TAILORED, ATS-optimized resume that:
1. Highlights relevant projects that match JD requirements
2. Uses exact keywords from the JD
3. Quantifies achievements where possible
4. Structures content for maximum ATS score

═══ OUTPUT JSON ═══
{{
    "resume": {{
        "professional_summary": "2-3 sentence summary tailored to this role...",
        "skills_section": ["Python", "React", "AWS", "...JD-matched skills"],
        "experience_section": [
            {{
                "title": "Project Lead / Developer",
                "organization": "Personal/Academic Project",
                "duration": "Estimated dates",
                "bullets": [
                    "Developed X using Y, achieving Z% improvement...",
                    "Led team of N to deliver..."
                ]
            }}
        ],
        "projects_section": [
            {{
                "name": "Project Name",
                "technologies": ["Tech1", "Tech2"],
                "description": "Brief ATS-optimized description",
                "achievements": ["Achievement 1", "Achievement 2"]
            }}
        ],
        "education_section": {{
            "degree": "Degree Name",
            "institution": "University",
            "year": "2024",
            "relevant_coursework": ["Course1", "Course2"]
        }}
    }},
    "suggestions": [
        {{
            "category": "experience",
            "suggestion": "Consider adding an internship or freelance work",
            "priority": "HIGH"
        }},
        {{
            "category": "skills",
            "suggestion": "Learn Kubernetes to match JD requirement",
            "priority": "MEDIUM"
        }}
    ],
    "skill_gaps": [
        {{
            "skill": "AWS",
            "importance": "HIGH",
            "learning_path": "AWS Cloud Practitioner certification (2-4 weeks)"
        }}
    ],
    "interview_prep_topics": [
        "Be ready to explain your REST API project architecture",
        "Prepare STAR stories for teamwork questions"
    ],
    "ats_optimization_notes": "This resume is optimized for X keywords. Include your GitHub link.",
    "estimated_ats_score": 82
}}
"""

    try:
        response = await engine.reason(ReasoningRequest(
            prompt=prompt,
            user_id=user_id,
            agent="campaign",
            mode=ReasoningMode.DEEP,
            max_thinking_tokens=32768,
            system_instruction="You are an expert resume writer. Output valid JSON only. Be specific and actionable."
        ))
        
        content = response.content
        if "```" in content:
            content = re.sub(r'```(?:json)?\s*', '', content).replace('```', '')
        
        match = re.search(r'\{[\s\S]*\}', content)
        if match:
            result = json.loads(match.group())
        else:
            result = json.loads(content)
        
        _create_thought_signature(
            user_id, "campaign",
            f"Generated resume for: {job_description[:100]}",
            f"Created tailored resume with {len(result.get('resume', {}).get('projects_section', []))} projects, ATS score: {result.get('estimated_ats_score', 'N/A')}",
            db_helper, context
        )
        
        return {"success": True, "data": result}
        
    except Exception as e:
        print(f"❌ Resume Generation Error: {e}")
        return {"success": False, "error": str(e)}


# =============================================================================
# SIMULACRUM: INTERVIEW DESIGNER
# =============================================================================
async def design_interview(db_helper, user_id, job_description: str, resume_text: Optional[str] = None,
                           interview_type: str = "technical", difficulty: str = "medium",
                           context=None) -> dict:
    """
    Design a comprehensive mock interview based on JD and resume.
    Returns structured interview prompt ready for Gemini Live.
    """
    engine = BicameralEngine()
    
    persona_configs = {
        "recruiter": {
            "name": "Sarah Chen",
            "role": "Technical Recruiter",
            "style": "Warm, professional, focused on culture fit and basic qualifications",
            "focus": ["motivation", "career goals", "team fit", "communication skills"]
        },
        "hiring_manager": {
            "name": "Michael Torres", 
            "role": "Engineering Manager",
            "style": "Direct, results-oriented, wants to see problem-solving and ownership",
            "focus": ["past achievements", "leadership", "conflict resolution", "project outcomes"]
        },
        "technical": {
            "name": "Dr. Anika Patel",
            "role": "Senior Engineer",
            "style": "Curious, probing, wants deep technical understanding",
            "focus": ["system design", "code quality", "debugging", "architecture decisions"]
        },
        "behavioral": {
            "name": "James Wilson",
            "role": "HR Director",
            "style": "Empathetic but thorough, uses STAR method",
            "focus": ["teamwork", "challenges overcome", "failures and learnings", "ethics"]
        }
    }
    
    persona = persona_configs.get(interview_type, persona_configs["technical"])
    
    difficulty_instructions = {
        "easy": "Be encouraging. Accept surface-level answers. Provide hints when the candidate struggles.",
        "medium": "Be professional. Ask 1-2 follow-up questions per answer. Probe for specifics.",
        "hard": "Be challenging. Always ask follow-ups. Challenge assumptions. Play devil's advocate.",
        "brutal": "Be skeptical. Interrupt if answers are vague. Ask rapid-fire follow-ups. Stress test."
    }
    
    prompt = f"""You are an expert Interview Designer for Kaironex.

═══ JOB DESCRIPTION ═══
{job_description[:5000]}

═══ CANDIDATE RESUME ═══
{resume_text[:5000] if resume_text else "Not provided - design general interview based on JD requirements"}

═══ CONFIGURATION ═══
Interview Type: {interview_type}
Difficulty: {difficulty}
Interviewer Persona: {persona['name']} ({persona['role']})
Style: {persona['style']}

═══ YOUR MISSION ═══
Design a comprehensive 20-30 minute mock interview.

1. **QUESTION BANK**: Create 8-12 questions covering:
   - Opening/Icebreaker (1-2)
   - Core Technical/Behavioral (4-6)
   - Situational/Problem-solving (2-3)
   - Closing/Questions for company (1)

2. **EVALUATION RUBRIC**: What makes a good answer for each question

3. **GEMINI LIVE PROMPT**: Create the exact system instruction for the live AI interviewer

═══ OUTPUT JSON ═══
{{
    "interview_id": "uuid",
    "duration_minutes": 25,
    "interviewer": {{
        "name": "{persona['name']}",
        "role": "{persona['role']}",
        "company": "TechCorp (Simulated)",
        "personality": "Professional but approachable"
    }},
    "question_bank": [
        {{
            "id": "q1",
            "category": "icebreaker",
            "question": "Tell me about yourself and what draws you to this role.",
            "follow_ups": ["What specific projects led you here?", "Why now?"],
            "good_answer_criteria": ["Shows passion", "Connects background to role", "Concise"],
            "red_flags": ["Too generic", "No mention of JD requirements"]
        }},
        {{
            "id": "q2",
            "category": "technical",
            "question": "Walk me through how you would design...",
            "follow_ups": ["How would you handle scale?", "What tradeoffs did you consider?"],
            "good_answer_criteria": ["Structured approach", "Mentions relevant tech"],
            "red_flags": ["Jumps to solution without clarifying", "No consideration of edge cases"]
        }}
    ],
    "gemini_live_prompt": "You are {persona['name']}, a {persona['role']} at TechCorp interviewing a candidate for [Role]. Your style is {persona['style']}...FULL INTERVIEW CONDUCTING INSTRUCTIONS...",
    "difficulty_modifiers": {{
        "current": "{difficulty}",
        "instruction": "{difficulty_instructions.get(difficulty, difficulty_instructions['medium'])}"
    }},
    "post_interview_rubric": {{
        "technical_competency": "1-10",
        "communication": "1-10",
        "problem_solving": "1-10",
        "culture_fit": "1-10",
        "overall": "1-10"
    }},
    "candidate_prep_notes": [
        "Be ready to discuss your most complex project in detail",
        "Prepare STAR stories for behavioral questions",
        "Have 2-3 thoughtful questions about the company/team"
    ]
}}
"""

    try:
        response = await engine.reason(ReasoningRequest(
            prompt=prompt,
            user_id=user_id,
            agent="campaign",
            mode=ReasoningMode.DEEP,
            max_thinking_tokens=24576,
            system_instruction="You are an expert interview designer. Create detailed, realistic interviews. Output valid JSON only."
        ))
        
        content = response.content
        if "```" in content:
            content = re.sub(r'```(?:json)?\s*', '', content).replace('```', '')
        
        match = re.search(r'\{[\s\S]*\}', content)
        if match:
            result = json.loads(match.group())
        else:
            result = json.loads(content)
        
        # Add interview ID if not present
        if "interview_id" not in result:
            result["interview_id"] = f"interview_{uuid.uuid4().hex[:8]}"
        
        # Store interview design for reference
        db_helper.update_campaign_state(user_id, {
            "last_interview_design": result,
            "interview_ready": True
        })
        
        _create_thought_signature(
            user_id, "campaign",
            f"Designed {interview_type} interview for {job_description[:50]}",
            f"Created {len(result.get('question_bank', []))} questions, {difficulty} difficulty",
            db_helper, context
        )
        
        return {"success": True, "data": result}
        
    except Exception as e:
        print(f"❌ Interview Design Error: {e}")
        return {"success": False, "error": str(e)}


# =============================================================================
# SKILL TREE: ENHANCED REASONING
# =============================================================================
async def generate_skill_tree(db_helper, user_id, target_role: str, current_skills: list,
                               education: Optional[str] = None, context=None) -> dict:
    """
    Generate a comprehensive skill tree with learning paths.
    """
    engine = BicameralEngine()
    
    prompt = f"""You are a Career Development Strategist designing a Skill Tree.

═══ CONTEXT ═══
Target Role: {target_role}
Current Skills: {', '.join(current_skills) if current_skills else 'Not specified'}
Education: {education or 'Not specified'}

═══ YOUR MISSION ═══
Create a visual skill tree (7-12 nodes) that shows the learning path from current state to job-ready.

RULES:
1. Start with Foundation skills (Level 1)
2. Build to Intermediate (Level 2) with prerequisites
3. End with Advanced/Specialized (Level 3)
4. Include both technical AND soft skills
5. Mark first 1-2 skills as "unlocked" (ready to start)
6. Add XP costs based on learning time (100 XP ≈ 1 week of focused learning)

═══ OUTPUT JSON ═══
{{
    "skill_tree": [
        {{
            "id": "skill_python_basics",
            "name": "Python Fundamentals",
            "level": 1,
            "status": "unlocked",
            "description": "Core Python syntax, data structures, OOP",
            "parent": null,
            "xpCost": 200,
            "icon": "🐍",
            "learning_resources": [
                {{"type": "course", "name": "Python for Everybody", "url": "coursera.org/..."}},
                {{"type": "practice", "name": "LeetCode Easy Problems", "count": 20}}
            ],
            "verification": "Complete 20 easy problems OR build a CLI tool"
        }},
        {{
            "id": "skill_web_dev",
            "name": "Web Development",
            "level": 2,
            "status": "locked",
            "description": "REST APIs, Flask/FastAPI",
            "parent": "skill_python_basics",
            "xpCost": 300,
            "icon": "🌐",
            "learning_resources": [...],
            "verification": "Build a REST API with 5+ endpoints"
        }}
    ],
    "recommended_path": ["skill_python_basics", "skill_web_dev", "skill_databases", "skill_cloud"],
    "time_to_job_ready": "3-4 months",
    "weekly_commitment": "15-20 hours",
    "milestones": [
        {{"week": 4, "title": "Fundamentals Complete", "skills": ["skill_python_basics"]}},
        {{"week": 8, "title": "Can Build Projects", "skills": ["skill_web_dev", "skill_databases"]}},
        {{"week": 12, "title": "Interview Ready", "skills": ["skill_cloud", "skill_system_design"]}}
    ]
}}
"""

    try:
        response = await engine.reason(ReasoningRequest(
            prompt=prompt,
            user_id=user_id,
            agent="campaign",
            mode=ReasoningMode.DEEP,
            max_thinking_tokens=16384,
            system_instruction="You are a career strategist. Create actionable skill trees. Output valid JSON only."
        ))
        
        content = response.content
        if "```" in content:
            content = re.sub(r'```(?:json)?\s*', '', content).replace('```', '')
        
        match = re.search(r'\{[\s\S]*\}', content)
        if match:
            result = json.loads(match.group())
        else:
            result = json.loads(content)
        
        # Persist to campaign state
        db_helper.update_campaign_state(user_id, {
            "skill_tree": result.get("skill_tree", []),
            "skill_tree_metadata": {
                "recommended_path": result.get("recommended_path", []),
                "time_to_job_ready": result.get("time_to_job_ready", "Unknown"),
                "milestones": result.get("milestones", [])
            }
        })
        
        _create_thought_signature(
            user_id, "campaign",
            f"Generated skill tree for {target_role}",
            f"Created {len(result.get('skill_tree', []))} skill nodes, estimated {result.get('time_to_job_ready', 'N/A')} to job-ready",
            db_helper, context
        )
        
        return {"success": True, "data": result}
        
    except Exception as e:
        print(f"❌ Skill Tree Generation Error: {e}")
        return {"success": False, "error": str(e)}


# =============================================================================
# MAIN ENTRY POINT: run_campaign_agent (Sync wrapper for Appwrite Functions)
# =============================================================================
def run_campaign_agent(db_helper, payload, context):
    """
    Career Strategist: Goals, Skill Tree, Resume, Interview
    Sync wrapper that delegates to async functions.
    """
    import asyncio
    
    user_id = payload.get('userId')
    if not user_id:
        return context.res.json({"error": "No userId"})

    try:
        context.log(f"⚔️ Campaign Brain processing for {user_id}")
        db_helper.log_heartbeat(user_id, f"EVENT:CAMPAIGN | {payload.get('type')}")
        
        event_type = payload.get('type', '')
        data = payload.get('data', {})
        
        # Route to appropriate handler
        if event_type == 'analyze_resume':
            # ATS Resume Analysis (supports text or PDF)
            resume_text = data.get('resumeText', '')
            resume_pdf = data.get('resumePdf', '')  # Base64 encoded PDF
            job_desc = data.get('jobDesc', '')
            
            # Extract text from PDF if provided
            if resume_pdf and not resume_text:
                resume_text = extract_text_from_pdf(resume_pdf)
                if not resume_text:
                    return context.res.json({"error": "Failed to extract text from PDF"}, 400)
            
            if not resume_text:
                return context.res.json({"error": "No resume text or PDF provided"}, 400)
            
            result = asyncio.run(analyze_resume_ats(db_helper, user_id, resume_text, job_desc, context))
            return context.res.json(result)
        
        elif event_type == 'generate_resume':
            # Resume Generator (returns JSON + optional DOCX)
            job_desc = data.get('jobDesc', '')
            projects = data.get('projects', [])
            skills = data.get('skills', [])
            education = data.get('education', '')
            experience = data.get('experience', '')
            include_docx = data.get('includeDocx', True)  # Generate DOCX by default
            
            if not job_desc:
                return context.res.json({"error": "No job description provided"}, 400)
            if not projects:
                return context.res.json({"error": "No projects provided"}, 400)
            
            result = asyncio.run(generate_tailored_resume(
                db_helper, user_id, job_desc, projects, skills, education, experience, context
            ))
            
            # Generate DOCX if requested and generation was successful
            if result.get('success') and include_docx:
                docx_base64 = generate_resume_docx(result.get('data', {}))
                if docx_base64:
                    result['docx'] = docx_base64
                    result['docx_filename'] = f"resume_{datetime.datetime.now().strftime('%Y%m%d_%H%M%S')}.docx"
            
            return context.res.json(result)
        
        elif event_type == 'design_interview':
            # Mock Interview Designer
            job_desc = data.get('jobDesc', '')
            resume_text = data.get('resumeText', '')
            interview_type = data.get('interviewType', 'technical')
            difficulty = data.get('difficulty', 'medium')
            
            if not job_desc:
                return context.res.json({"error": "No job description provided"}, 400)
            
            result = asyncio.run(design_interview(
                db_helper, user_id, job_desc, resume_text, interview_type, difficulty, context
            ))
            return context.res.json(result)
        
        elif event_type == 'generate_skill_tree':
            # Skill Tree Generation
            target_role = data.get('targetRole', 'Software Engineer')
            current_skills = data.get('skills', [])
            education = data.get('education', '')
            
            result = asyncio.run(generate_skill_tree(
                db_helper, user_id, target_role, current_skills, education, context
            ))
            return context.res.json(result)
        
        elif event_type == 'campaign_calibration':
            # Full Campaign Calibration (Skill Tree + Quest Board + Armory init)
            target_role = data.get('targetRole', 'Software Engineer')
            current_skills = data.get('skills', [])
            education = data.get('education', '')
            
            # Generate skill tree
            skill_result = asyncio.run(generate_skill_tree(
                db_helper, user_id, target_role, current_skills, education, context
            ))
            
            if not skill_result.get('success'):
                return context.res.json(skill_result)
            
            # Initialize armory with default items
            armory_data = {
                "inventory": [
                    {"id": "item_resume_builder", "name": "Resume Forge", "type": "TOOL", "status": "equipped", "costXp": 0},
                    {"id": "item_interview_prep", "name": "Mock Interview", "type": "TOOL", "status": "owned", "costXp": 0}
                ],
                "blueprints": [
                    {"id": "bp_portfolio", "name": "Portfolio Builder", "costXp": 500},
                    {"id": "bp_linkedin", "name": "LinkedIn Optimizer", "costXp": 300}
                ]
            }
            
            # Initialize quest board with guidance quests
            quest_board = [
                {"id": "quest_skill_1", "title": "Master First Skill", "type": "SKILL", "description": "Complete the first skill in your tree", "xp": 200, "status": "active", "tags": ["Foundation"]},
                {"id": "quest_resume", "title": "Forge Your Resume", "type": "RESEARCH", "description": "Use the Armory to create your first ATS-optimized resume", "xp": 300, "status": "active", "tags": ["Preparation"]},
                {"id": "quest_mock", "title": "Face the Simulacrum", "type": "NETWORKING", "description": "Complete one mock interview session", "xp": 400, "status": "active", "tags": ["Practice"]}
            ]
            
            # Save to campaign state
            db_helper.update_campaign_state(user_id, {
                "armory": armory_data,
                "quest_board": quest_board,
                "is_calibrated": True
            })
            
            return context.res.json({
                "success": True,
                "data": {
                    "skill_tree": skill_result.get("data", {}),
                    "armory": armory_data,
                    "quest_board": quest_board
                }
            })
        
        elif event_type == 'new_goal':
            # Legacy goal setting
            goal_title = payload.get('goal_title', 'New Goal')
            ai = GeminiClient()
            prompt = f"User set a career goal: {goal_title}. Provide 3 strategic first steps."
            strategy = ai.generate_response(prompt)
            
            _create_thought_signature(user_id, "campaign", prompt, strategy, db_helper, context)
            
            db_helper.create_intervention(user_id, "NEW_GOAL", strategy, strategy="PROACTIVE")
            db_helper.update_state_cache(user_id, {"campaign": {"active_quest": goal_title, "current_strategy": strategy}})
            
            return context.res.json({"status": "goal_set", "strategy": strategy})
        
        elif event_type == 'scan_jobs':
            # Job scan (uses search grounding - may hit quota)
            success = scan_daily_jobs(db_helper, user_id, context)
            return context.res.json({"status": "scan_complete" if success else "scan_skipped"})
        
        else:
            return context.res.json({"status": "unknown_event", "type": event_type})

    except Exception as e:
        context.error(f"❌ Campaign Agent Error: {e}")
        return context.res.json({"error": str(e)}, 500)


# =============================================================================
# DAILY JOB SCAN (Search Grounding - Uses Quota)
# =============================================================================
def scan_daily_jobs(db_helper, user_id, context):
    """
    Daily Job Search using Google Search Grounding.
    Note: Uses free tier quota - may be limited.
    """
    try:
        # Get Campaign State
        campaign_doc = db_helper.db.get_document(
            database_id=APPWRITE_DATABASE_ID, 
            collection_id=CAMPAIGN_STATE_COL, 
            document_id=user_id
        )
        
        # Check last scan time
        last_scan_str = campaign_doc.get('last_job_scan')
        if last_scan_str:
            last_scan = datetime.datetime.fromisoformat(last_scan_str.replace("Z", "+00:00"))
            if last_scan.tzinfo is None:
                last_scan = last_scan.replace(tzinfo=datetime.timezone.utc)
            
            now = datetime.datetime.now(datetime.timezone.utc)
            if (now - last_scan).total_seconds() < 86400:
                context.log(f"⏳ Job scan skipped for {user_id}. Last scan: {last_scan_str}")
                return False

        # Get target role from profile
        user_doc = db_helper.get_user_doc(user_id)
        profile_json = user_doc.get('studentprofile_json', '{}')
        profile = json.loads(profile_json)
        
        target_role = profile.get('targetRole')
        location = profile.get('currentCountry', 'Remote')
        
        if not target_role:
            context.log(f"⚠️ Job scan skipped: No target role set.")
            return False

        context.log(f"🔎 Scanning jobs for {user_id}: {target_role} in {location}")

        ai = GeminiClient()
        prompt = f"""Find 3 REAL, active job listings for a '{target_role}' role in '{location}' (or Remote).
        Focus on entry-level/internships.
        Return JSON array: [{{"id": "uuid", "title": "Role @ Company", "type": "APPLICATION", "description": "Apply via...", "xp": 500, "status": "active", "tags": ["Remote"]}}]"""
        
        try:
            raw_response = ai.generate_response(prompt, use_search=True, model_type="flash")
            job_json_str = raw_response.strip().replace("```json", "").replace("```", "")
            
            if "quota" in job_json_str.lower():
                raise Exception("Quota exceeded")
                
            new_quests = json.loads(job_json_str)
            
        except Exception as e:
            if "quota" in str(e).lower():
                context.log(f"⚠️ Search quota exceeded")
                new_quests = [{"id": f"sys_quota_{uuid.uuid4().hex[:4]}", "title": "Job Search Paused (Quota)", "type": "APPLICATION", "description": "Will retry tomorrow.", "xp": 0, "status": "active", "tags": ["System"]}]
            else:
                context.log(f"❌ Job scan error: {e}")
                return False
        
        # Update quest board
        current_board = json.loads(campaign_doc.get('quest_board_json', '[]'))
        updated_board = new_quests + current_board[:10]
        
        db_helper.db.update_document(
            database_id=APPWRITE_DATABASE_ID,
            collection_id=CAMPAIGN_STATE_COL,
            document_id=user_id,
            data={
                "quest_board_json": json.dumps(updated_board),
                "last_job_scan": datetime.datetime.now(datetime.timezone.utc).isoformat()
            }
        )
        
        _create_thought_signature(user_id, "campaign", f"Job scan for {target_role}", f"Found {len(new_quests)} listings", db_helper, context)
        return True

    except Exception as e:
        context.log(f"❌ Daily Job Scan Error: {e}")
        return False
